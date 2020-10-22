# -*- coding: utf-8 -*-

import os
import sys
import csv
import itertools
from openpyxl import Workbook
from openpyxl.chart import Reference, Series, ScatterChart
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.styles import numbers
from datetime import datetime

# 「設定」
DATA_SHEET_NAME = "Data"
CHART_SHEET_NAME = "Graph"
CHART_TITLE = "The Waveform of The Triangle Wave"
Y_AXIS_TITLE = "Voltage V(dBV)"
X_AXIS_TITLE = "Frequency f(Hz)"
LEGEND_POSITION = "r"
RANGE_TO_ADD_CHART = "B3"
EXTENTION_OF_FILE = ".xlsx"
BEGIN_DATA_ROW = 1
X_AXIS_COLUMN = 1
CHART_HEIGHT = 12
CHART_WIDTH = 24
Y_AXIS_MIN = -2.50E+02
Y_AXIS_MAX = 0.00E+00
X_AXIS_MIN = 0.00E+00
X_AXIS_MAX = 9.75E+03

# ----------------ここからはプログラム-----------------------------

def convert_type(datas):

    convert_datas = []

    for data in datas:
        try:
            # 数値データはfloatに変換
            convert_datas.append(float(data))

        except ValueError:

            try:
                # 日付データはdatetimeに変換
                data_timestamp = datetime.strptime(data, "%Y/%m/%d %H:%M:%S")
                convert_datas.append(data_timestamp)

            except ValueError:
                convert_datas.append(data)
                pass

    return convert_datas

def main():

    # コマンドライン引数取得
    args = sys.argv
    csv_file_path = args[1]
    y_axis_columns = []
    x_axis_column = X_AXIS_COLUMN

    for count in range(2, len(args)):
        y_axis_columns.append(args[count])

    # 出力先ファイルパスの設定
    current_time = datetime.now()
    timestamp = current_time.strftime("%Y%m%d-%H%M%S")
    split_csv_file_name = os.path.splitext(csv_file_path)
    output_file_path = "excel_graphs/" + split_csv_file_name[0] + "_" + timestamp + EXTENTION_OF_FILE

    # ワークブック、ワークシートオブジェクトを作成
    work_book = Workbook()
    data_sheet = work_book.active
    data_sheet.title = DATA_SHEET_NAME
    chart_sheet = work_book.create_sheet(title=CHART_SHEET_NAME)
    work_book.active = chart_sheet

    # csvファイルの読み込み
    convert_datas = []
    with open(csv_file_path, encoding='utf-8') as csv_file:
        csv_reader = csv.reader(csv_file, delimiter=',', quotechar='"')

        for line in csv_reader:
            convert_datas = convert_type(line)
            data_sheet.append(convert_datas)

        csv_columns = len(line)

        # CSVファイルの列数分、列番号をリストに追加
        if y_axis_columns == []:
            y_axis_columns = range(x_axis_column + 1, csv_columns + 1)

    # グラフの書式設定
    chart = ScatterChart()
    chart.title = CHART_TITLE
    chart.y_axis.title = Y_AXIS_TITLE
    chart.x_axis.title = X_AXIS_TITLE
    chart.y_axis.scaling.min = Y_AXIS_MIN
    chart.y_axis.scaling.max = Y_AXIS_MAX
    chart.x_axis.scaling.min = X_AXIS_MIN
    chart.x_axis.scaling.max = X_AXIS_MAX
    # chart.legend.position = LEGEND_POSITION
    chart.height = CHART_HEIGHT
    chart.width = CHART_WIDTH

    # X軸の参照先の設定
    line_number = csv_reader.line_num
    x_axis_values = Reference(data_sheet, min_col=x_axis_column, min_row=BEGIN_DATA_ROW + 1, max_row=line_number)

    # Y軸の参照先の設定
    for y_axis_column in y_axis_columns:
        y_axis_values = Reference(data_sheet, min_col=y_axis_column, min_row=BEGIN_DATA_ROW, max_row=line_number)
        series = Series(y_axis_values, x_axis_values, title_from_data=True)
        chart.series.append(series)

    # グラフの挿入
    chart_sheet.add_chart(chart, RANGE_TO_ADD_CHART)
    work_book.save(output_file_path)

    print("Result file is [%s]" % output_file_path)

if __name__ == '__main__':
    main()